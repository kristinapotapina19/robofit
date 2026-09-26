"""
Параметры объекта: проверка, шаблон и загрузка из файла (п. 3.2.2–3.2.5 ТЗ).

Шаблон повторяет формат демо-датасета организатора (Параметр | Ед. изм. | Значение |
Мин | Макс | Пояснение), поэтому загрузить можно и заполненный шаблон, и сам файл
организатора, и CSV с разделителем «;» или «,».
"""
from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def to_number(v: Any) -> float | None:
    if _is_number(v):
        return float(v)
    if isinstance(v, str):
        s = v.replace(" ", "").replace(" ", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def validate(preset: dict, values: dict, required: set[str]) -> dict:
    """Проверка типов, обязательности и диапазонов. Сообщения — с подсказкой, как исправить."""
    errors, warnings, clean = [], [], {}
    for p in preset["parameters"]:
        label = p["label"]
        raw = values.get(label, p["default"])
        numeric = _is_number(p["default"])
        if raw is None or raw == "":
            if label in required:
                errors.append({"label": label, "message": "Обязательный параметр не заполнен",
                               "fix": f"Укажите значение в {p['unit'] or 'единицах параметра'}"
                                      + (f", например {p['default']}" if p["default"] is not None else "")})
            clean[label] = None
            continue
        if numeric:
            v = to_number(raw)
            if v is None:
                errors.append({"label": label, "message": f"Ожидается число, получено «{raw}»",
                               "fix": "Введите число без единиц измерения, дробную часть — через точку или запятую"})
                clean[label] = raw
                continue
            mn, mx = to_number(p.get("min")), to_number(p.get("max"))
            if v < 0 and (mn is None or mn >= 0):
                errors.append({"label": label, "message": "Значение не может быть отрицательным",
                               "fix": "Введите неотрицательное число"})
            if mn is not None and mx is not None and not (mn <= v <= mx) and mn != mx:
                warnings.append({"label": label,
                                 "message": f"Значение {v:g} вне типового диапазона {mn:g}–{mx:g} {p['unit'] or ''}",
                                 "fix": "Проверьте единицы измерения; если значение верное — расчёт выполнится, "
                                        "но результат выходит за область, на которой проверялась модель"})
            clean[label] = v
        else:
            clean[label] = raw
    unknown = [k for k in values if k not in {p["label"] for p in preset["parameters"]}]
    for k in unknown:
        warnings.append({"label": k, "message": "Параметр не относится к выбранному типу объекта и не используется",
                         "fix": "Проверьте, что выбран правильный тип объекта"})
    return {"ok": not errors, "errors": errors, "warnings": warnings, "values": clean}


def template_xlsx(preset: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = preset["title"][:31]
    ws.append([f"ШАБЛОН ПАРАМЕТРОВ ОБЪЕКТА: {preset['title']}. Заполните столбец «Значение»."])
    ws["A1"].font = Font(bold=True)
    ws.append(["Параметр", "Ед. изм.", "Значение", "Мин", "Макс", "Пояснение / источник"])
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EDE7F6")
    section = None
    for p in preset["parameters"]:
        if p.get("section") and p["section"] != section:
            section = p["section"]
            ws.append([f"▌{section}"])
        ws.append([p["label"], p["unit"], p["default"], p["min"], p["max"], p["source_note"]])
    for col, w in zip("ABCDEF", (60, 14, 18, 10, 10, 60)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_csv(preset: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Параметр", "Ед. изм.", "Значение"])
    for p in preset["parameters"]:
        w.writerow([p["label"], p["unit"], p["default"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _rows_from_upload(filename: str, content: bytes, sheet_hint: str | None) -> list[list[Any]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb[sheet_hint] if sheet_hint and sheet_hint in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith((".csv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        delim = ";" if text.count(";") >= text.count(",") else ","
        return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]
    raise ValueError("Поддерживаются файлы .xlsx и .csv — сохраните файл в одном из этих форматов")


def parse_upload(filename: str, content: bytes, preset: dict) -> dict:
    rows = _rows_from_upload(filename, content, preset["title"])
    labels = {p["label"].strip().lower(): p["label"] for p in preset["parameters"]}
    value_col = 2
    for row in rows[:5]:
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if "значение" in cells:
            value_col = cells.index("значение")
            break
    values, matched = {}, 0
    for row in rows:
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip().lower()
        if key in labels:
            v = row[value_col] if len(row) > value_col else None
            if len(row) == 2:          # формат «параметр;значение»
                v = row[1]
            values[labels[key]] = v
            matched += 1
    if not matched:
        raise ValueError("В файле не найдено ни одного параметра выбранного типа объекта. "
                         "Скачайте шаблон и заполните столбец «Значение», не меняя названия параметров")
    return {"values": values, "matched": matched, "total": len(preset["parameters"])}
