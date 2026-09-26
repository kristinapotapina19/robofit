"""
Выгрузка результатов (п. 3.7 ТЗ).

Excel — расчётные таблицы: параметры объекта, подбор, сценарии, трассировка формул,
реестр допущений с источниками.
PDF — отчёт для печати и отправки руководству.

В обоих форматах обязательна пометка, что результат является предварительной
оценкой и требует верификации при обследовании объекта (п. 3.7.5).
"""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DISCLAIMER = ("Результат является предварительной оценкой на основе демонстрационных "
              "данных и допущений, указанных в отчёте. Перед принятием инвестиционного "
              "решения требуется обследование объекта и подтверждение характеристик "
              "у поставщиков решений.")

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "C:/Windows/Fonts/arial.ttf",
]

HEAD_FILL = PatternFill("solid", fgColor="EDE7F6")
TITLE_FONT = Font(bold=True, size=13)
HEAD_FONT = Font(bold=True)
THIN = Side(style="thin", color="D0CCC4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ==========================================================================
# Excel
# ==========================================================================
def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list],
           widths: list[int] | None = None, note: str | None = None) -> None:
    ws = wb.create_sheet(title[:31])
    r = 1
    if note:
        ws.cell(r, 1, note).font = Font(italic=True, color="666666")
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(len(headers), 2))
        ws.row_dimensions[r].height = 30
        r += 2
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(r, c, h)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.border = BORDER
    for row in rows:
        r += 1
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=isinstance(value, str) and len(str(value)) > 40,
                                       vertical="top")
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=(4 if note else 2), column=1)


def build_xlsx(payload: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    summary = wb.create_sheet("Сводка")
    summary["A1"] = "Отчёт по оценке роботизации"
    summary["A1"].font = Font(bold=True, size=15)
    meta = [
        ("Объект", payload["object"]["title"]),
        ("Операция", payload["operation"]["title"]),
        ("Решение", payload["solution"]["name"]),
        ("Поставщик", payload["solution"]["vendor"]),
        ("Парк, ед.", payload["fleet"]["count"]),
        ("Основание парка", payload["fleet"]["basis"]),
        ("Рекомендация", payload["recommendation"]["verdict"]),
        ("Интерпретация", payload["recommendation"]["interpretation"]),
        ("Дата расчёта", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Версия каталога / модели", f"{payload.get('catalog_version', '')} / "
                                     f"{payload.get('model_version', '')}"),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        summary.cell(i, 1, k).font = HEAD_FONT
        summary.cell(i, 2, v)
    row = len(meta) + 5
    summary.cell(row, 1, "Показатель").font = HEAD_FONT
    for c, sc in enumerate(payload["scenarios"], start=2):
        summary.cell(row, c, sc["title"]).font = HEAD_FONT
    metrics = [
        ("Роботов, шт.", "robot_count"),
        ("CAPEX, руб.", "capex_total"), ("OPEX, руб./год", "opex_year"),
        ("Изменение OPEX к базовому, руб./год", "opex_change"),
        ("Экономия ФОТ, руб./год", "labor_saving_year"),
        ("Годовой эффект, руб.", "annual_effect"),
        ("Эффект с учётом амортизации, руб./год", "annual_effect_after_depreciation"),
        ("Окупаемость, лет", "payback_years"), ("ROI, %", "roi_percent"),
        ("TCO за горизонт, руб.", "tco"),
    ]
    for i, (label, key) in enumerate(metrics, start=row + 1):
        summary.cell(i, 1, label).font = HEAD_FONT
        for c, sc in enumerate(payload["scenarios"], start=2):
            summary.cell(i, c, sc.get(key))
    summary.column_dimensions["A"].width = 38
    for c in range(2, 2 + len(payload["scenarios"])):
        summary.column_dimensions[get_column_letter(c)].width = 20
    last = row + len(metrics) + 2
    summary.cell(last, 1, DISCLAIMER).alignment = Alignment(wrap_text=True)
    summary.merge_cells(start_row=last, start_column=1, end_row=last, end_column=4)
    summary.row_dimensions[last].height = 46

    _sheet(wb, "Параметры объекта", ["Параметр", "Ед. изм.", "Значение", "Источник"],
           [[p["label"], p.get("unit"), p.get("value"), p.get("source")]
            for p in payload["object"]["parameters"]],
           [46, 14, 18, 60],
           note=payload["object"].get("data_source"))

    _sheet(wb, "Подбор решений", ["Решение", "Поставщик", "Цена, руб.", "Вердикт",
                                  "Ранг", "Причины", "Ограничения", "Нет данных"],
           [[c["name"], c["vendor"], c["price_rub"], c["verdict"], c["score"],
             "; ".join(c["reasons"]), "; ".join(c["warnings"]) or c.get("decisive_reason") or "",
             ", ".join(c["missing_specs"])] for c in payload["candidates"]],
           [44, 26, 14, 16, 8, 60, 50, 24],
           note="Показаны все рассмотренные кандидаты с причиной включения или исключения.")

    trace_rows = [["Общие шаги", st["name"], st["formula"], st["result"], st["unit"],
                   "; ".join(f"{k}: {v}" for k, v in st["inputs"].items())]
                  for st in payload.get("steps", [])]
    for sc in payload["scenarios"]:
        for st in sc.get("steps", []):
            trace_rows.append([sc["title"], st["name"], st["formula"],
                               st["result"], st["unit"],
                               "; ".join(f"{k}: {v}" for k, v in st["inputs"].items())])
    _sheet(wb, "Трассировка расчёта",
           ["Сценарий", "Шаг", "Формула", "Результат", "Ед. изм.", "Входные данные"],
           trace_rows, [24, 34, 52, 18, 14, 70])

    _sheet(wb, "Допущения", ["Допущение", "Значение", "Ед. изм.", "Источник"],
           [[k, v["value"], v["unit"], v["source"]]
            for k, v in payload["assumptions"].items()],
           [34, 14, 34, 70],
           note="Недокументированные коэффициенты в модели не используются (п. 3.5.1 ТЗ).")

    if payload.get("sensitivity"):
        names = {"purchase": "Покупка", "raas": "RaaS"}
        _sheet(wb, "Чувствительность", ["Параметр", "Изменение, %", "Сценарий",
                                         "Годовой эффект, руб.", "Изменение эффекта, руб.",
                                         "Окупаемость, лет", "TCO, руб.", "Смена интервала"],
               [[r["factor"], r["delta_percent"], names.get(r["scenario"], r["scenario"]),
                 r["annual_effect"], r["annual_effect_change"], r["payback_years"], r["tco"],
                 "да" if r["crosses_band"] else ""] for r in payload["sensitivity"]],
               [26, 14, 12, 20, 22, 16, 18, 16],
               note="Чувствительность к трём параметрам (п. 3.5.6 ТЗ): стоимость оборудования, "
                    "объём операций, стоимость труда.")
    _sheet(wb, "Риски и ограничения", ["№", "Риск / ограничение"],
           [[i, r] for i, r in enumerate(dict.fromkeys(payload.get("risks", []) + payload.get("warnings", [])), 1)],
           [6, 120], note=DISCLAIMER)

    sim = payload.get("simulation")
    if sim:
        _sheet(wb, "Имитация", ["Показатель", "Значение"],
               [["Парк, ед.", payload["fleet"]["count"]],
                ["Цикл рейса, с", sim["cycle_time_s"]],
                ["Производительность робота по имитации, оп./ч", sim["theoretical_per_robot_h"]],
                ["Требуется, оп./ч", sim["required_per_h"]],
                ["Закрыто, оп./ч", sim["achieved_per_h"]],
                ["Загрузка парка", sim["utilization"]],
                ["Доля простоя", sim["idle_share"]],
                ["Доля зарядки", sim["charging_share"]],
                ["Среднее ожидание задачи, с", sim["avg_wait_s"]],
                ["Узкое место", sim["bottleneck"]],
                ["Вывод", sim["verdict"]]],
               [56, 60])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ==========================================================================
# PDF
# ==========================================================================
def _register_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    for path in FONT_CANDIDATES:
        if Path(path).exists():
            pdfmetrics.registerFont(TTFont("Report", path))
            bold = path.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf").replace(
                "arial.ttf", "arialbd.ttf")
            if Path(bold).exists():
                pdfmetrics.registerFont(TTFont("Report-Bold", bold))
                registerFontFamily("Report", normal="Report", bold="Report-Bold",
                                   italic="Report", boldItalic="Report-Bold")
            else:
                registerFontFamily("Report", normal="Report", bold="Report",
                                   italic="Report", boldItalic="Report")
            return "Report"
    return "Helvetica"          # кириллица не отобразится, но отчёт соберётся


def build_pdf(payload: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    font = _register_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName=font,
                          fontSize=9.5, leading=13)
    h1 = ParagraphStyle("h1", parent=body, fontSize=17, leading=21, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=body, fontSize=12.5, leading=16,
                        spaceBefore=14, spaceAfter=6, textColor=colors.HexColor("#4B2E83"))
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11,
                           textColor=colors.HexColor("#666666"))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title="Отчёт по оценке роботизации",
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm)

    def wrap(value):
        """Длинный текст оборачивается в Paragraph, иначе reportlab его обрежет."""
        if isinstance(value, str) and len(value) > 55:
            return Paragraph(value, ParagraphStyle("cell", parent=body, fontSize=8.4,
                                                   leading=10.5))
        return value

    def table(data, widths, align_right=()):
        data = [[wrap(c) for c in row] for row in data]
        t = Table(data, colWidths=widths, repeatRows=1)
        style = [
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 8.6),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDE7F6")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5D1C8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for c in align_right:
            style.append(("ALIGN", (c, 1), (c, -1), "RIGHT"))
        t.setStyle(TableStyle(style))
        return t

    story = [Paragraph("Отчёт по оценке роботизации", h1)]
    story.append(Paragraph(
        f"Объект: {payload['object']['title']}. Операция: {payload['operation']['title']}. "
        f"Дата расчёта: {datetime.now().strftime('%d.%m.%Y')}.", body))

    story.append(Paragraph("Выбранное решение", h2))
    sol = payload["solution"]
    story.append(table([
        ["Наименование", sol["name"]],
        ["Поставщик", sol["vendor"]],
        ["Цена единицы, руб.", f"{sol['price_rub']:,.0f}".replace(",", " ")],
        ["Парк, ед.", str(payload["fleet"]["count"])],
        ["Основание парка", payload["fleet"]["basis"]],
    ], [45 * mm, 125 * mm]))

    story.append(Paragraph("Сравнение сценариев", h2))
    short = {"baseline": "Без роботизации", "purchase": "Покупка", "raas": "RaaS"}
    header = ["Показатель"] + [short.get(s["code"], s["title"]) for s in payload["scenarios"]]
    metrics = [
        ("CAPEX, млн руб.", "capex_total", 1e6),
        ("OPEX, млн руб./год", "opex_year", 1e6),
        ("Годовой эффект, млн руб.", "annual_effect", 1e6),
        ("Окупаемость, лет", "payback_years", 1),
        ("ROI, %", "roi_percent", 1),
        ("TCO за горизонт, млн руб.", "tco", 1e6),
    ]
    rows = [header]
    for label, key, scale in metrics:
        row = [label]
        for sc in payload["scenarios"]:
            v = sc.get(key)
            row.append("—" if v is None else f"{v / scale:,.1f}".replace(",", " "))
        rows.append(row)
    width = (170 * mm - 55 * mm) / max(len(payload["scenarios"]), 1)
    story.append(table(rows, [55 * mm] + [width] * len(payload["scenarios"]),
                       align_right=tuple(range(1, len(payload["scenarios"]) + 1))))

    rec = payload.get("recommendation")
    if rec:
        story.append(Paragraph("Заключение", h2))
        story.append(Paragraph(f"<b>{rec['verdict']}.</b> {rec['interpretation']}", body))

    risks = list(dict.fromkeys(payload.get("risks", []) + payload.get("warnings", [])))
    if risks:
        story.append(Paragraph("Риски и ограничения", h2))
        for r in risks:
            story.append(Paragraph(f"• {r}", body))

    sens = [r for r in payload.get("sensitivity", []) if r["scenario"] == "purchase"]
    if sens:
        story.append(Paragraph("Чувствительность (сценарий покупки)", h2))
        rows = [["Параметр", "Изменение", "Годовой эффект, млн руб.", "Окупаемость, лет"]]
        for r in sens:
            rows.append([r["factor"], f"{r['delta_percent']:+d}%",
                         f"{r['annual_effect'] / 1e6:,.1f}".replace(",", " "),
                         "—" if r["payback_years"] is None else f"{r['payback_years']:.2f}"])
        story.append(table(rows, [55 * mm, 25 * mm, 45 * mm, 45 * mm], align_right=(2, 3)))

    sim = payload.get("simulation")
    if sim:
        story.append(Paragraph("Проверка имитацией", h2))
        story.append(Paragraph(sim["verdict"], body))
        story.append(Spacer(1, 6))
        story.append(table([
            ["Цикл рейса, с", f"{sim['cycle_time_s']:.0f}"],
            ["Производительность робота по имитации, оп./ч",
             f"{sim['theoretical_per_robot_h']:.1f}"],
            ["Требуется / закрыто, оп./ч",
             f"{sim['required_per_h']:.1f} / {sim['achieved_per_h']:.1f}"],
            ["Загрузка парка", f"{sim['utilization'] * 100:.0f}%"],
            ["Узкое место", sim["bottleneck"]],
        ], [80 * mm, 90 * mm]))

    story.append(PageBreak())
    story.append(Paragraph("Как посчитано", h2))
    purchase = next((s for s in payload["scenarios"] if s["code"] == "purchase"),
                    payload["scenarios"][-1])
    rows = [["Шаг", "Формула", "Результат"]]
    for st in payload.get("steps", []) + purchase.get("steps", []):
        res = st["result"]
        res = f"{res:,.2f}".replace(",", " ") if isinstance(res, (int, float)) else str(res)
        rows.append([st["name"], st["formula"], f"{res} {st['unit']}"])
    story.append(table(rows, [45 * mm, 90 * mm, 35 * mm]))

    story.append(Paragraph("Допущения и источники", h2))
    rows = [["Допущение", "Значение", "Источник"]]
    for key, v in payload["assumptions"].items():
        rows.append([key, f"{v['value']}", v["source"]])
    story.append(table(rows, [42 * mm, 22 * mm, 106 * mm]))

    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Версия каталога {payload.get('catalog_version', '')}, "
                           f"версия расчётной модели {payload.get('model_version', '')}.", small))
    story.append(Paragraph(DISCLAIMER, small))

    doc.build(story)
    return buffer.getvalue()
